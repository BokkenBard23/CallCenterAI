"""Export router — Excel and PDF export of analysis results.

Endpoints:
  GET /export/{session_id}/excel — export analysis as .xlsx
  GET /export/{session_id}/pdf    — export analysis as .pdf
"""

from __future__ import annotations

import io
import logging
from datetime import datetime
from typing import Optional

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse

from app.utils.session import session_store

logger = logging.getLogger(__name__)

router = APIRouter()


@router.get(
    "/export/{session_id}/excel",
    summary="Export analysis results as Excel (.xlsx)",
)
async def export_excel(session_id: str) -> StreamingResponse:
    """Export analysis results for a session as an Excel workbook.

    The workbook contains 3 sheets:
      - «Диалог»: turn_index, speaker, text
      - «Совпадения»: phrase_text, matched_text, cascade_order, speaker, etc.
      - «Сводка»: LLM summary text (if available)
    """
    session = session_store.get(session_id)
    if session is None:
        raise HTTPException(
            status_code=404,
            detail=f"Session '{session_id}' not found.",
        )

    if not session.analyses:
        raise HTTPException(
            status_code=400,
            detail="No analysis results available for this session.",
        )

    try:
        import openpyxl
    except ImportError:
        raise HTTPException(
            status_code=503,
            detail="Excel export unavailable: openpyxl not installed.",
        )

    wb = openpyxl.Workbook()

    # ── Sheet 1: Диалог ───────────────────────────────────────
    ws_dialog = wb.active
    ws_dialog.title = "Диалог"
    ws_dialog.append(["#", "Спикер", "Текст"])

    if session.dialog and session.dialog.turns:
        for turn in session.dialog.turns:
            ws_dialog.append([turn.turn_index, turn.speaker, turn.text])

    # ── Sheet 2: Совпадения ──────────────────────────────────
    ws_matches = wb.create_sheet("Совпадения")
    ws_matches.append([
        "Фраза из словаря",
        "Совпавший текст",
        "Порядок каскада",
        "Спикер",
        "Тип совпадения",
        "Дистанция (слова)",
        "Точное совпадение",
        "Ограничение канала",
    ])

    for analysis in session.analyses.values():
        if analysis.search_result and analysis.search_result.matches:
            for match in analysis.search_result.matches:
                ws_matches.append([
                    match.phrase_text,
                    match.matched_text,
                    match.cascade_order,
                    match.speaker,
                    match.match_type,
                    match.word_distance,
                    "Да" if match.is_exact_match else "Нет",
                    match.channel_constraint,
                ])

    # ── Sheet 3: Сводка ──────────────────────────────────────
    ws_summary = wb.create_sheet("Сводка")
    ws_summary.append(["Параметр", "Значение"])

    for analysis in session.analyses.values():
        if analysis.llm_result and analysis.llm_result.summary:
            ws_summary.append(["Резюме", analysis.llm_result.summary])
        if analysis.llm_result and analysis.llm_result.topic:
            ws_summary.append(["Тема", analysis.llm_result.topic])
        if analysis.llm_result and analysis.llm_result.result:
            ws_summary.append(["Результат", analysis.llm_result.result])
        if analysis.llm_result and analysis.llm_result.client_sentiment:
            ws_summary.append(["Настроение клиента", analysis.llm_result.client_sentiment])
        if analysis.llm_result and analysis.llm_result.resolution:
            ws_summary.append(["Разрешение", analysis.llm_result.resolution])
        if analysis.llm_result and analysis.llm_result.key_points:
            for idx, point in enumerate(analysis.llm_result.key_points, 1):
                ws_summary.append([f"Ключевой момент {idx}", point])

    # ── Stream response ──────────────────────────────────────
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    date_str = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"analysis_{session_id}_{date_str}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get(
    "/export/{session_id}/pdf",
    summary="Export analysis results as PDF",
)
async def export_pdf(session_id: str) -> StreamingResponse:
    """Export analysis results for a session as a PDF document.

    The PDF contains:
      - Title: «Результаты анализа диалога»
      - Section 1: Dialogue turns
      - Section 2: Matches table
      - Section 3: LLM summary (if available)
    """
    session = session_store.get(session_id)
    if session is None:
        raise HTTPException(
            status_code=404,
            detail=f"Session '{session_id}' not found.",
        )

    if not session.analyses:
        raise HTTPException(
            status_code=400,
            detail="No analysis results available for this session.",
        )

    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import cm, mm
        from reportlab.platypus import (
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError:
        raise HTTPException(
            status_code=503,
            detail="PDF export unavailable: reportlab not installed.",
        )

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CustomTitle",
        parent=styles["Title"],
        fontSize=16,
        spaceAfter=20,
    )
    heading_style = ParagraphStyle(
        "CustomHeading",
        parent=styles["Heading2"],
        fontSize=12,
        spaceAfter=10,
        spaceBefore=15,
    )
    body_style = styles["Normal"]

    elements: list = []

    # ── Title ─────────────────────────────────────────────────
    elements.append(Paragraph("Результаты анализа диалога", title_style))
    elements.append(Spacer(1, 10))

    # ── Section 1: Dialogue turns ─────────────────────────────
    elements.append(Paragraph("Диалог", heading_style))

    if session.dialog and session.dialog.turns:
        dialog_data = [["#", "Спикер", "Текст"]]
        for turn in session.dialog.turns:
            # Truncate long texts for table display
            text = turn.text if len(turn.text) <= 200 else turn.text[:197] + "..."
            dialog_data.append([str(turn.turn_index), turn.speaker, text])

        dialog_table = Table(dialog_data, colWidths=[1.5 * cm, 3 * cm, 13 * cm])
        dialog_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
            ("ALIGN", (0, 0), (0, -1), "CENTER"),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 10),
            ("FONTSIZE", (0, 1), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.lightgrey]),
        ]))
        elements.append(dialog_table)
    else:
        elements.append(Paragraph("Диалог не загружен", body_style))

    elements.append(Spacer(1, 15))

    # ── Section 2: Matches table ─────────────────────────────
    elements.append(Paragraph("Совпадения", heading_style))

    all_matches: list = []
    for analysis in session.analyses.values():
        if analysis.search_result and analysis.search_result.matches:
            all_matches.extend(analysis.search_result.matches)

    if all_matches:
        matches_data = [["Фраза", "Уровень", "Спикер", "Тип"]]
        for match in all_matches:
            phrase = match.phrase_text if len(match.phrase_text) <= 50 else match.phrase_text[:47] + "..."
            matches_data.append([
                phrase,
                str(match.cascade_order),
                match.speaker,
                "Точное" if match.is_exact_match else match.match_type,
            ])

        matches_table = Table(matches_data, colWidths=[8 * cm, 2 * cm, 4 * cm, 4 * cm])
        matches_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 10),
            ("FONTSIZE", (0, 1), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.lightgrey]),
        ]))
        elements.append(matches_table)
    else:
        elements.append(Paragraph("Совпадений не найдено", body_style))

    elements.append(Spacer(1, 15))

    # ── Section 3: LLM Summary ──────────────────────────────
    for analysis in session.analyses.values():
        if analysis.llm_result and analysis.llm_result.summary:
            elements.append(Paragraph("Сводка", heading_style))
            elements.append(Paragraph(f"<b>Резюме:</b> {analysis.llm_result.summary}", body_style))
            if analysis.llm_result.topic:
                elements.append(Paragraph(f"<b>Тема:</b> {analysis.llm_result.topic}", body_style))
            if analysis.llm_result.result:
                elements.append(Paragraph(f"<b>Результат:</b> {analysis.llm_result.result}", body_style))
            if analysis.llm_result.client_sentiment:
                elements.append(Paragraph(f"<b>Настроение клиента:</b> {analysis.llm_result.client_sentiment}", body_style))
            if analysis.llm_result.resolution:
                elements.append(Paragraph(f"<b>Разрешение:</b> {analysis.llm_result.resolution}", body_style))
            if analysis.llm_result.key_points:
                points_text = "<br/>".join(
                    f"• {point}" for point in analysis.llm_result.key_points
                )
                elements.append(Paragraph(f"<b>Ключевые моменты:</b><br/>{points_text}", body_style))
            break  # Only include the first analysis with LLM result

    # ── Build PDF ─────────────────────────────────────────────
    doc.build(elements)
    buffer.seek(0)

    date_str = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"analysis_{session_id}_{date_str}.pdf"

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
